from pathlib import Path
import sys
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'build' / 'report_tools'))
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = ROOT / 'laporan'
DOCX = OUT / 'Laporan_Catatan_Harian_Muhammad_Alif_Mujaddid.docx'
GREEN = '315E4B'
MUTED = '59665E'
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(1.9)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.3)
section.header_distance = Cm(.9)
section.footer_distance = Cm(.9)
section.different_first_page_header_footer = True
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string('25332B')
normal.paragraph_format.line_spacing = 1.16
normal.paragraph_format.space_after = Pt(8)
for name, size in [('Title', 32), ('Heading 1', 19), ('Heading 2', 13), ('Heading 3', 11)]:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(GREEN)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(8)
for name in ['Caption', 'Subtitle']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor.from_string(MUTED)
doc.styles['Caption'].font.size = Pt(9)

header = section.header.paragraphs[0]
header.text = 'CATATAN HARIAN  |  LAPORAN PENGEMBANGAN APLIKASI'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('Muhammad Alif Mujaddid  •  240605110082     |     ')
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string(MUTED)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)


list_counter = 0

def p(text='', style=None):
    global list_counter
    if style == 'List Number':
        list_counter += 1
        para = doc.add_paragraph(f'{list_counter}. {text}')
        para.paragraph_format.left_indent = Cm(.55)
        para.paragraph_format.first_line_indent = Cm(-.55)
        return para
    return doc.add_paragraph(text, style)


def h(text, level=1):
    global list_counter
    list_counter = 0
    return doc.add_heading(text, level)


def page(title):
    doc.add_page_break()
    h(title)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    props.append(shd)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, text in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = text
        shade(c, GREEN)
        for r in c.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
    for index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if index % 2 == 0:
                shade(cells[i], 'F0F4EE')
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        trPr.append(cant)
        for i, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Cm(widths[i])
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(4)
                para.paragraph_format.space_after = Pt(5)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.size = Pt(9.5)
    p().paragraph_format.space_after = Pt(0)
    return t


def code(text):
    para = p()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(10)
    para.paragraph_format.line_spacing = 1.05
    for line in text.splitlines():
        run = para.add_run(line + '\n')
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
    props = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'F0F4EE')
    props.append(shd)


def shots(items):
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (filename, caption) in enumerate(items):
        t.columns[i].width = Cm(8)
        c = t.cell(0, i)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        para = c.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(4)
        para.add_run().add_picture(str(OUT/'screenshots'/filename), width=Cm(6.65))
        cap = t.cell(1, i).paragraphs[0]
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        r = cap.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:'+side)
        border.set(qn('w:val'), 'nil')
        borders.append(border)
    t._tbl.tblPr.append(borders)
    p().paragraph_format.space_after = Pt(0)

# 1 — Sampul
p('LAPORAN PENGEMBANGAN APLIKASI').runs[0].font.color.rgb = RGBColor.from_string(GREEN)
p('UTS • MULTIPLATFORM', 'Subtitle')
p()
p('CATATAN\nHARIAN', 'Title')
p('Aplikasi pencatatan lokal berbasis Flutter\ndengan SQLite dan Shared Preferences', 'Subtitle')
p()
identity = [
    ('Nama', 'Muhammad Alif Mujaddid'),
    ('NIM', '240605110082'),
    ('Kelas', '[Isi kelas]'),
    ('Dosen pengampu', '[Isi nama dosen pengampu]'),
]
table(['IDENTITAS', 'KETERANGAN'], identity, [4.3, 11.9])
p()
p('UIN MAULANA MALIK IBRAHIM MALANG').runs[0].bold = True
p('2026')
p('Tanggal penyusunan: 23 Oktober 2026')
p('Dokumen menjelaskan implementasi, rancangan database, tampilan aplikasi, dan hasil pengujian berdasarkan kode proyek yang tersedia.').runs[0].italic = True

# 2 — Ringkasan dan isi
page('RINGKASAN DAN DAFTAR ISI')
h('Ringkasan', 2)
p('Catatan Harian merupakan aplikasi untuk menulis, menyimpan, dan mengelola catatan pribadi secara lokal. Aplikasi menyediakan operasi tambah, baca, ubah, dan hapus; pencarian berdasarkan judul atau isi; kategori; sematan; serta warna kartu. Antarmuka dibangun dengan Flutter dan Material 3, sedangkan data catatan disimpan menggunakan SQLite.')
p('Pengembangan juga mencakup tema terang dan gelap, tata letak grid/daftar, penghitung kata, validasi isian, konfirmasi penghapusan, serta konfirmasi ketika pengguna meninggalkan perubahan yang belum disimpan. Tema dan pilihan urutan disimpan melalui Shared Preferences.')
p('Pemeriksaan proyek mencakup 14 tes otomatis yang lulus, analisis statis tanpa masalah, dan APK debug Android yang telah berhasil dibuat. Delapan screenshot di laporan dirender dari widget aplikasi sebenarnya dengan data demonstrasi.')
h('Daftar isi', 2)
table(['BAGIAN', 'HALAMAN'], [
    ('Sampul', '1'),
    ('Ringkasan dan daftar isi', '2'),
    ('BAB I — Pendahuluan', '3'),
    ('BAB II — Perancangan aplikasi', '4'),
    ('BAB III — Perancangan database', '5'),
    ('BAB IV — Implementasi dan alur program', '6'),
    ('BAB V — Dokumentasi tampilan: beranda', '7'),
    ('Dokumentasi tampilan: editor dan pengaturan', '8'),
    ('Dokumentasi tampilan: pencarian dan filter', '9'),
    ('Dokumentasi tampilan: validasi dan konfirmasi', '10'),
    ('BAB VI — Hasil pengujian', '11'),
    ('BAB VII — Penggunaan, kesimpulan, dan pengembangan', '12'),
], [13.7, 2.5])

# 3
page('BAB I — PENDAHULUAN')
h('1.1 Latar belakang', 2)
p('Kegiatan belajar dan aktivitas sehari-hari menghasilkan banyak informasi singkat, seperti ide proyek, jadwal persiapan ujian, dan rencana kegiatan. Informasi tersebut perlu dicatat agar dapat dibaca kembali. Ketika jumlah catatan bertambah, pengguna juga memerlukan pencarian dan pengelompokan yang mudah dipahami.')
p('Catatan Harian dikembangkan sebagai aplikasi pencatatan dengan penyimpanan lokal. Pendekatan ini memungkinkan pengguna mengelola catatan tanpa akun dan tanpa koneksi internet. Flutter digunakan untuk membangun antarmuka dari satu basis kode, dengan fokus verifikasi build pada Android.')
h('1.2 Rumusan masalah', 2)
for text in [
    'Bagaimana membuat aplikasi yang dapat menambah, menampilkan, mengubah, dan menghapus catatan?',
    'Bagaimana menyimpan catatan dan preferensi pengguna agar tetap tersedia ketika aplikasi dibuka kembali?',
    'Bagaimana membantu pengguna menemukan dan mengatur catatan melalui pencarian, kategori, dan sematan?',
    'Bagaimana menyajikan antarmuka yang nyaman serta menangani isian tidak valid dan kegagalan operasi?',
]: p(text, 'List Number')
h('1.3 Tujuan dan manfaat', 2)
p('Tujuan pengembangan adalah menghasilkan aplikasi catatan yang mudah digunakan, memiliki penyimpanan terstruktur, dan mendukung pengaturan koleksi. Dari sisi pembelajaran, proyek menerapkan widget Flutter, pengelolaan state melalui StatefulWidget dan setState, navigasi antarlayar, operasi asinkron, SQLite, serta penyimpanan key-value.')
h('1.4 Batasan ruang lingkup', 2)
p('Catatan terdiri atas judul, isi, waktu pembuatan, kategori, warna, dan status sematan. Kategori yang tersedia adalah Pribadi, Kuliah, Kerja, dan Ide. Aplikasi belum memiliki akun pengguna, sinkronisasi cloud, pengingat, lampiran, maupun ekspor/impor cadangan. Penyimpanan aplikasi desktop dan web belum dikonfigurasi.')

# 4
page('BAB II — PERANCANGAN APLIKASI')
h('2.1 Kebutuhan fungsional', 2)
table(['KEBUTUHAN', 'PERILAKU APLIKASI'], [
    ('Pengelolaan catatan', 'Menyimpan judul dan isi, membuka detail untuk diedit, serta menghapus setelah konfirmasi.'),
    ('Pencarian dan filter', 'Mencocokkan judul/isi tanpa membedakan kapital; menggabungkan kategori dan status sematan.'),
    ('Pengaturan koleksi', 'Menyematkan catatan, memilih urutan waktu, dan beralih grid/daftar.'),
    ('Personalisasi', 'Memilih empat warna kartu dan mengganti tema terang/gelap.'),
    ('Umpan balik', 'Menampilkan loading, validasi, pesan kegagalan, dan status simpan.'),
], [4.2, 12])
h('2.2 Arsitektur dan pembagian tanggung jawab', 2)
code('main.dart — CatatanApp dan tema\n  |\n  +-- HomeScreen <---- navigasi ----> NoteFormScreen\n  |       |                              |\n  |       +------ DatabaseHelper --------+\n  |                    |\n  |                 SQLite\n  +-- SettingsService --> Shared Preferences')
p('HomeScreen menangani koleksi, pencarian, filter, dan aksi kartu. NoteFormScreen menangani editor, validasi, serta pengaturan catatan. DatabaseHelper memusatkan akses database, sedangkan model Note mengubah objek Dart menjadi Map dan sebaliknya. AppTheme menyatukan warna serta gaya komponen.')
h('2.3 Teknologi yang digunakan', 2)
p('Lingkungan verifikasi menggunakan Flutter 3.35.5 dan Dart 3.9.2. Dependensi utama adalah sqflite ^2.4.2, path ^1.9.1, dan shared_preferences ^2.5.5. Pengujian menggunakan flutter_test dan sqflite_common_ffi ^2.3.7+1. Batas versi tersebut mengikuti pubspec.yaml; versi hasil resolusi tercatat di pubspec.lock.')

# 5
page('BAB III — PERANCANGAN DATABASE')
h('3.1 Penyimpanan terstruktur', 2)
p('Database lokal bernama catatan.db dibuat otomatis melalui DatabaseHelper. Versi skema saat ini adalah 2. Satu tabel notes digunakan karena seluruh atribut catatan dapat direpresentasikan sebagai satu entitas tanpa relasi antartabel.')
table(['KOLOM', 'TIPE', 'FUNGSI / KETENTUAN'], [
    ('id', 'INTEGER', 'Primary key, AUTOINCREMENT; identitas unik catatan.'),
    ('title', 'TEXT', 'Judul catatan; NOT NULL.'),
    ('content', 'TEXT', 'Isi catatan; NOT NULL.'),
    ('created_at', 'TEXT', 'Waktu pembuatan dalam format ISO 8601; NOT NULL.'),
    ('category', 'TEXT', "Kategori; nilai awal 'Pribadi'; NOT NULL."),
    ('color_index', 'INTEGER', 'Indeks palet warna; nilai awal 0; NOT NULL.'),
    ('is_pinned', 'INTEGER', '1 berarti disematkan, 0 berarti tidak; nilai awal 0.'),
], [3.1, 2.6, 10.5])
h('3.2 Operasi CRUD', 2)
p('Insert menambahkan data baru tanpa menyertakan id agar SQLite membuat identitas otomatis. GetNotes membaca data dengan urutan sematan terlebih dahulu, kemudian waktu pembuatan. Update dan delete menggunakan klausa where id = ? dengan argumen terpisah. Ketika catatan diedit, createdAt tetap menggunakan waktu pembuatan awal.')
h('3.3 Migrasi versi 1 ke versi 2', 2)
p('Migrasi menggunakan onUpgrade untuk menambahkan kolom category, color_index, dan is_pinned melalui ALTER TABLE. Nilai awal memungkinkan catatan lama dibaca tanpa perlu mengisi atribut baru secara manual. Tidak ada penghapusan tabel dalam migrasi tersebut. Tes migrasi memeriksa judul, isi, nilai awal, dan urutan sematan setelah pembaruan skema.')
h('3.4 Preferensi pengguna', 2)
p('Shared Preferences menyimpan dua nilai boolean: dark_mode untuk tema dan sort_newest untuk urutan. Isi catatan tidak disimpan di Shared Preferences. Pilihan grid/daftar dan filter saat ini merupakan state antarmuka selama layar digunakan.')

# 6
page('BAB IV — IMPLEMENTASI DAN ALUR PROGRAM')
h('4.1 Alur membuka aplikasi', 2)
p('main() menginisialisasi binding Flutter lalu menjalankan CatatanApp. Aplikasi membaca preferensi tema. HomeScreen membaca preferensi urutan dan memuat catatan secara asinkron. Selama proses berlangsung, antarmuka menampilkan indikator loading. Jika pemuatan gagal, pesan kesalahan dan tombol Coba lagi ditampilkan.')
h('4.2 Alur menambah dan mengedit catatan', 2)
for text in [
    'Pengguna menekan Tulis catatan atau mengetuk kartu yang sudah ada.',
    'Editor menyiapkan TextEditingController untuk judul dan isi. Jumlah kata dihitung dari isi yang dipisah berdasarkan whitespace.',
    'Pengguna dapat membuka Atur catatan untuk memilih kategori, warna kartu, dan status sematan.',
    'Saat Simpan ditekan, Form memeriksa bahwa judul dan isi tidak kosong setelah trim.',
    'Objek Note dibuat, lalu operasi insert atau update dijalankan. Tombol simpan dinonaktifkan selama proses agar tidak ditekan berulang.',
    'Jika berhasil, editor mengembalikan hasil ke beranda. Koleksi dimuat ulang dan pesan keberhasilan ditampilkan.',
]: p(text, 'List Number')
h('4.3 Pencarian, filter, dan urutan', 2)
p('Judul dan isi diubah ke huruf kecil sebelum dibandingkan dengan kata kunci yang sudah di-trim. Hasil tersebut digabungkan dengan syarat kategori dan status sematan. Catatan yang disematkan selalu ditempatkan lebih dahulu; urutan terbaru/terlama berlaku pada masing-masing kelompok. Pilihan tampilan menyesuaikan jumlah kolom dengan lebar layar.')
h('4.4 Perlindungan dan penanganan kegagalan', 2)
p('PopScope mencegah pengguna meninggalkan editor yang memiliki perubahan tanpa konfirmasi. Penghapusan juga memerlukan persetujuan pada dialog. Operasi memuat, menyimpan, menghapus, dan mengubah sematan memiliki batas waktu 15 detik. Batas waktu memberi umpan balik ketika operasi terlalu lama, tetapi tidak menjamin pembatalan operasi database yang sudah berjalan.')
p('Jika penyimpanan gagal, editor tetap terbuka dan tombol simpan diaktifkan kembali. Pesan timeout penyimpanan mengarahkan pengguna untuk memeriksa daftar catatan. Implementasi detail tersedia di lib/screens/note_form_screen.dart dan lib/screens/home_screen.dart.')

# 7
page('BAB V — DOKUMENTASI TAMPILAN')
h('5.1 Beranda dan tema aplikasi', 2)
p('Beranda menampilkan identitas aplikasi, tanggal, ringkasan koleksi, pencarian, filter, serta kartu catatan. Ikon bulan/matahari digunakan untuk mengganti tema.')
shots([
    ('01_beranda_terang.png', 'Gambar 1. Beranda dalam tema terang.'),
    ('02_beranda_gelap.png', 'Gambar 2. Beranda dalam tema gelap.'),
])
p('Kartu yang disematkan ditandai ikon pin dan berada di bagian atas. Tombol Tulis catatan membuka editor. Tombol panah mengubah urutan, sedangkan tombol tata letak beralih antara grid dan daftar.')
p('Metode screenshot: render widget Flutter asli melalui flutter test, viewport 420 × 860 piksel logis dan hasil PNG 840 × 1720 piksel. Data demonstrasi terdiri atas tiga catatan. Gambar bukan tangkapan dari perangkat Android fisik dan tidak menyertakan bilah sistem operasi.').runs[0].italic = True

# 8
page('DOKUMENTASI — EDITOR DAN PENGATURAN')
h('5.2 Menulis dan mempersonalisasi catatan', 2)
p('Editor menyediakan judul, isi, penghitung kata, serta tombol Simpan. Panel Atur catatan dibuka melalui ikon pengaturan di kanan atas.')
shots([
    ('03_editor_catatan.png', 'Gambar 3. Editor dengan data catatan contoh.'),
    ('04_pengaturan_catatan.png', 'Gambar 4. Kategori, warna, dan sematan.'),
])
p('Pada contoh, catatan berkategori Kuliah dan disematkan. Pengguna dapat memilih Pribadi, Kuliah, Kerja, atau Ide, mengganti warna kartu, serta mengubah status sematan. Tombol Selesai hanya menutup panel; tombol Simpan pada editor menyimpan perubahan ke database.')
p('Judul dan isi wajib diisi. Editor menggunakan area isi yang dapat menampung beberapa baris, sedangkan tata letaknya dapat digulir saat tinggi layar terbatas atau keyboard terbuka.')

# 9
page('DOKUMENTASI — PENCARIAN DAN FILTER')
h('5.3 Menemukan catatan dengan lebih cepat', 2)
p('Pencarian dan filter diproses dari koleksi catatan yang telah dimuat. Filter kategori dapat digunakan bersamaan dengan filter sematan dan kata kunci.')
shots([
    ('05_pencarian.png', 'Gambar 5. Pencarian dengan kata kunci UTS.'),
    ('06_filter_kategori_sematan.png', 'Gambar 6. Filter Kuliah dan Disematkan.'),
])
p('Gambar 5 menunjukkan bahwa kata UTS menemukan catatan Persiapan UTS Multiplatform. Tombol silang pada kolom pencarian membersihkan kata kunci. Gambar 6 memperlihatkan satu hasil yang memenuhi dua syarat: kategori Kuliah dan status disematkan.')
p('Jika tidak ada hasil, aplikasi menampilkan keadaan kosong beserta pilihan Reset filter. Pengguna dapat menggeser deretan kategori secara horizontal untuk membuka kategori lain pada layar kecil.')

# 10
page('DOKUMENTASI — VALIDASI DAN KONFIRMASI')
h('5.4 Mencegah input kosong dan penghapusan tidak sengaja', 2)
p('Aplikasi memberi penjelasan langsung pada isian yang belum lengkap dan meminta konfirmasi sebelum menghapus catatan.')
shots([
    ('07_validasi_form.png', 'Gambar 7. Validasi judul dan isi yang kosong.'),
    ('08_konfirmasi_hapus.png', 'Gambar 8. Dialog konfirmasi penghapusan.'),
])
p('Gambar 7 diambil setelah tombol Simpan ditekan tanpa mengisi form. Pesan Judul wajib diisi dan Isi wajib diisi muncul, sehingga penyimpanan tidak dilanjutkan. Gambar 8 menampilkan nama catatan yang akan dihapus serta pilihan Batal atau Hapus.')
p('Pada screenshot konfirmasi, penghapusan belum dijalankan. Pengujian CRUD database dilakukan secara terpisah. Selain dialog ini, aplikasi juga meminta konfirmasi ketika pengguna meninggalkan editor dengan perubahan yang belum disimpan.')

# 11
page('BAB VI — HASIL PENGUJIAN')
h('6.1 Metode dan hasil', 2)
p('Pengujian dilakukan melalui flutter_test untuk interaksi widget serta sqflite_common_ffi untuk database pengujian. Analisis statis dijalankan dengan flutter analyze --no-pub. Pemeriksaan ulang pada 23 Oktober 2026 menghasilkan 14 tes lulus dan analisis tanpa masalah.')
table(['KELOMPOK', 'SKENARIO', 'HASIL'], [
    ('Database (3 tes)', 'CRUD; berbagi inisialisasi database; migrasi versi 1 serta urutan sematan.', '3/3 lulus'),
    ('Beranda (4 tes)', 'Filter kategori + sematan di layar kecil; pencarian dan reset; gagal memuat lalu coba lagi; timeout pemuatan.', '4/4 lulus'),
    ('Editor (7 tes)', 'Simpan kategori/warna/sematan; layar kecil + keyboard; konfirmasi perubahan; validasi; gagal simpan; timeout simpan; simpan sukses dan kembali.', '7/7 lulus'),
    ('Analisis statis', 'Pemeriksaan Dart/Flutter pada proyek.', 'Tanpa masalah'),
    ('Build Android', 'APK debug sudah berhasil dibuat pada pengerjaan aplikasi.', 'Berhasil'),
], [3.4, 10.1, 2.7])
h('6.2 Bukti dan cara reproduksi', 2)
code('flutter analyze --no-pub\nflutter test --no-pub\nflutter build apk --debug --no-pub')
p('Bukti teks pemeriksaan ulang disertakan sebagai laporan/hasil_analisis.txt dan laporan/hasil_pengujian.txt. Pada akhir log pengujian tercantum All tests passed!. Pengambilan delapan screenshot dijalankan terpisah dari 14 tes fungsional menggunakan tools/report_capture_test.dart.')
h('6.3 Batas verifikasi', 2)
p('Tes widget memverifikasi antarmuka melalui lingkungan pengujian Flutter, bukan pengoperasian perangkat fisik. Keberhasilan build menunjukkan APK dapat dihasilkan; laporan ini tidak mengklaim pengujian instalasi dan pemakaian pada Android fisik. Build iOS/macOS serta penyimpanan pada web, Windows, dan Linux belum diverifikasi.')

# 12
page('BAB VII — PENGGUNAAN DAN KESIMPULAN')
h('7.1 Menjalankan dan membuat APK', 2)
p('Siapkan Flutter SDK, Android SDK, serta emulator atau perangkat Android. Buka terminal pada direktori proyek yang berisi pubspec.yaml, kemudian gunakan perintah berikut:')
code('flutter pub get\nflutter devices\nflutter run -d <device-id>\nflutter build apk --debug')
p('Gunakan ID dari flutter devices untuk mengganti <device-id>. APK debug tersedia di build/app/outputs/flutter-apk/app-debug.apk. Gunakan instalasi Flutter yang sama untuk mengambil dependensi, menjalankan tes, dan membuat build. Konfigurasi release saat ini masih memakai signing debug sehingga memerlukan konfigurasi signing sendiri sebelum publikasi.')
h('7.2 Kesimpulan', 2)
p('Catatan Harian telah menerapkan pengelolaan catatan lokal dengan Flutter, SQLite, dan Shared Preferences. Pencarian, kategori, sematan, pilihan warna, dan tema membantu pengguna mengelola koleksi. Validasi serta dialog konfirmasi memberi perlindungan terhadap isian kosong dan tindakan yang tidak disengaja. Hasil 14 tes yang lulus serta analisis bersih mendukung kesesuaian perilaku pada skenario yang diuji.')
h('7.3 Saran pengembangan', 2)
p('Pengembangan berikutnya dapat menambahkan ekspor/impor cadangan, sinkronisasi opsional, pengingat, atau lampiran. Dukungan database web dan desktop memerlukan backend tambahan. Pengujian perangkat fisik serta aksesibilitas dengan ukuran teks besar dapat dilakukan sebelum distribusi yang lebih luas.')
h('7.4 Sumber dokumentasi proyek', 2)
for text in [
    'README.md dan pubspec.yaml: fitur, dependensi, petunjuk menjalankan, serta batas platform.',
    'lib/main.dart dan lib/theme/app_theme.dart: inisialisasi aplikasi dan tema.',
    'lib/models/note.dart dan lib/db/database_helper.dart: model, skema, migrasi, serta operasi data.',
    'lib/screens/*.dart dan lib/services/settings_service.dart: antarmuka, validasi, navigasi, dan preferensi.',
    'test/*.dart: skenario pengujian fungsional; tools/report_capture_test.dart: pengambilan screenshot.',
]: p(text, 'List Bullet')

props = doc.core_properties
props.title = 'Laporan Pengembangan Aplikasi Catatan Harian'
props.subject = 'UTS Multiplatform — Flutter, SQLite, dan Shared Preferences'
props.author = 'Muhammad Alif Mujaddid'
props.keywords = '240605110082, Catatan Harian, Flutter, UTS, Multiplatform'
props.comments = 'Screenshot dirender dari widget Flutter dengan data demonstrasi. Kelas dan dosen dapat dilengkapi pada sampul.'
doc.save(DOCX)
print(DOCX)
print('embedded_images:', len(doc.inline_shapes))
